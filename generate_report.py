import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import Counter
import statistics

file_path = r"C:\Users\ThinkPad\Documents\antigravity\wise-newton\Điều_tra_Sốt_xuất_huyết_Thực_hành_cộng_đồng_-_all_versions_-_labels_-_2026-09-18-02-12-28_1.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
headers = rows[0]
data = rows[1:]
N = len(data)

def col(name):
    idx = headers.index(name)
    return [r[idx] for r in data]

def get_multi_stat(prefix, label_cleaner=None):
    sub_cols = [h for h in headers if h.startswith(prefix) and '/' in h]
    results = []
    for sc in sub_cols:
        idx = headers.index(sc)
        cnt = sum(1 for r in data if r[idx] == 1)
        sub_label = sc.split('/')[-1].strip()
        if label_cleaner:
            sub_label = label_cleaner(sub_label)
        results.append((sub_label, cnt, cnt / N * 100))
    results.sort(key=lambda x: x[1], reverse=True)
    return results

# Styles
out_wb = openpyxl.Workbook()
default_sheet = out_wb.active

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
title_font = Font(name="Calibri", size=12, bold=True, color="1F4E79")
section_font = Font(name="Calibri", size=11, bold=True, color="002060")
section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

def add_table_sheet(sheet_title, tables):
    ws_out = out_wb.create_sheet(title=sheet_title)
    ws_out.views.sheetView[0].showGridLines = True
    current_row = 1
    
    for t_title, headers_list, rows_list in tables:
        ws_out.cell(row=current_row, column=1, value=t_title).font = title_font
        current_row += 1
        
        # Header
        for col_idx, h in enumerate(headers_list, 1):
            cell = ws_out.cell(row=current_row, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        current_row += 1
        
        # Rows
        for r in rows_list:
            is_section = len(r) > 1 and r[1] == ""
            for col_idx, val in enumerate(r, 1):
                cell = ws_out.cell(row=current_row, column=col_idx, value=val)
                cell.border = thin_border
                if is_section:
                    cell.font = section_font
                    cell.fill = section_fill
                else:
                    cell.font = Font(name="Calibri", size=10)
                    if col_idx > 1:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1
        current_row += 2

    for col_cells in ws_out.columns:
        max_len = max(len(str(c.value or "")) for c in col_cells)
        col_letter = openpyxl.utils.get_column_letter(col_cells[0].column)
        ws_out.column_dimensions[col_letter].width = max(max_len + 4, 14)

# ==================== TABLE 1: DÂN SỐ - XÃ HỘI ====================
ages_raw = col("2. Tuổi")
cleaned_ages = []
age_groups = []
for r in data:
    v = r[headers.index("2. Tuổi")]
    if v is not None:
        try:
            val = float(str(v).strip())
            if val > 1900: val = 2026 - val
            elif val == 1868: val = 2026 - 1968
            cleaned_ages.append(val)
            if val < 30: age_groups.append("< 30 tuổi")
            elif val < 50: age_groups.append("30 - 49 tuổi")
            elif val < 65: age_groups.append("50 - 64 tuổi")
            else: age_groups.append("≥ 65 tuổi")
        except: age_groups.append("Khuyết")
    else: age_groups.append("Khuyết")

ag_c = Counter(age_groups)
g_c = Counter(col("3. Giới tính"))
eth_c = Counter(col("4. Dân tộc"))
edu_c = Counter(col("5. Trình độ học vấn"))
job_c = Counter(col("6. Nghề nghiệp"))

t1_rows = [
    ["1. Nhóm tuổi", "", ""],
    ["  < 30 tuổi", ag_c["< 30 tuổi"], f"{ag_c['< 30 tuổi']/N*100:.1f}%"],
    ["  30 - 49 tuổi", ag_c["30 - 49 tuổi"], f"{ag_c['30 - 49 tuổi']/N*100:.1f}%"],
    ["  50 - 64 tuổi", ag_c["50 - 64 tuổi"], f"{ag_c['50 - 64 tuổi']/N*100:.1f}%"],
    ["  ≥ 65 tuổi", ag_c["≥ 65 tuổi"], f"{ag_c['≥ 65 tuổi']/N*100:.1f}%"],
    ["  Khuyết (Missing)", ag_c["Khuyết"], f"{ag_c['Khuyết']/N*100:.1f}%"],
    ["  Tuổi trung bình ± ĐLC (Min - Max)", f"{statistics.mean(cleaned_ages):.1f} ± {statistics.stdev(cleaned_ages):.1f}", f"Trung vị: {statistics.median(cleaned_ages):.1f} ({min(cleaned_ages):.0f} - {max(cleaned_ages):.0f})"],
    ["2. Giới tính", "", ""],
    ["  Nữ", g_c["Nữ"], f"{g_c['Nữ']/N*100:.1f}%"],
    ["  Nam", g_c["Nam"], f"{g_c['Nam']/N*100:.1f}%"],
    ["  Khuyết", g_c[None], f"{g_c[None]/N*100:.1f}%"],
    ["3. Dân tộc", "", ""],
    ["  Kinh", eth_c["Kinh"], f"{eth_c['Kinh']/N*100:.1f}%"],
    ["  Khơ-me", eth_c["Khơ-me"], f"{eth_c['Khơ-me']/N*100:.1f}%"],
    ["  Hoa", eth_c["Hoa"], f"{eth_c['Hoa']/N*100:.1f}%"],
    ["  Khuyết", eth_c[None], f"{eth_c[None]/N*100:.1f}%"],
    ["4. Trình độ học vấn", "", ""],
    ["  Không biết chữ / Dưới tiểu học", edu_c["Không biết chữ/dưới tiểu học"], f"{edu_c['Không biết chữ/dưới tiểu học']/N*100:.1f}%"],
    ["  Tiểu học", edu_c["Tiểu học"], f"{edu_c['Tiểu học']/N*100:.1f}%"],
    ["  Trung học cơ sở", edu_c["Trung học cơ sở"], f"{edu_c['Trung học cơ sở']/N*100:.1f}%"],
    ["  Trung học phổ thông", edu_c["Trung học phổ thông"], f"{edu_c['Trung học phổ thông']/N*100:.1f}%"],
    ["  Trung cấp, Cao đẳng, Đại học", edu_c["Trung cấp, Cao đẳng, Đại học"], f"{edu_c['Trung cấp, Cao đẳng, Đại học']/N*100:.1f}%"],
    ["  Sau đại học", edu_c["Sau đại học"], f"{edu_c['Sau đại học']/N*100:.1f}%"],
    ["  Khuyết", edu_c[None], f"{edu_c[None]/N*100:.1f}%"],
    ["5. Nghề nghiệp", "", ""],
    ["  Làm ruộng / Làm vườn", job_c["Làm ruộng/làm vườn"], f"{job_c['Làm ruộng/làm vườn']/N*100:.1f}%"],
    ["  Nội trợ", job_c["Nội trợ"], f"{job_c['Nội trợ']/N*100:.1f}%"],
    ["  Buôn bán", job_c["Buôn bán"], f"{job_c['Buôn bán']/N*100:.1f}%"],
    ["  Công việc văn phòng", job_c["Công việc văn phòng"], f"{job_c['Công việc văn phòng']/N*100:.1f}%"],
    ["  Khác (Lao động tự do, hưu trí...)", job_c["Khác (điền mục 6.1)"], f"{job_c['Khác (điền mục 6.1)']/N*100:.1f}%"],
    ["  Khuyết", job_c[None], f"{job_c[None]/N*100:.1f}%"]
]

add_table_sheet("1. Dân số - Xã hội", [
    ("BẢNG 1: ĐẶC ĐIỂM DÂN SỐ - XÃ HỘI CỦA ĐỐI TƯỢNG KHẢO SÁT (N = 601)", ["Đặc tính", "Tần số (n)", "Tỷ lệ (%)"], t1_rows)
])

# ==================== TABLE 2: KIẾN THỨC VỀ SXH ====================
q7_c = Counter(col("7. Anh/chị đã từng nghe về bệnh SXH? "))
q9_c = Counter(col("9. Bệnh SXH lây như thế nào? "))
q10_c = Counter(col("10. Nếu do muỗi, Anh/chị có biết muỗi gì gây SXH?"))
q11_c = Counter(col("11. Muỗi gây bệnh SXH thường chích vào lúc nào?"))

t2_rows = [
    ["1. Đã từng nghe về bệnh Sốt xuất huyết (C7)", "", ""],
    ["  Có", q7_c["Có"], f"{q7_c['Có']/N*100:.1f}%"],
    ["  Không", q7_c["Không"], f"{q7_c['Không']/N*100:.1f}%"],
    ["  Khuyết", q7_c[None], f"{q7_c[None]/N*100:.1f}%"],
    ["2. Nguồn thông tin tiếp nhận về SXH (C8 - Đa lựa chọn)", "", ""]]

for label, cnt, pct in get_multi_stat("8."):
    if "Anh/chị nghe thông tin" not in label:
        t2_rows.append([f"  {label}", cnt, f"{pct:.1f}%"])

t2_rows.extend([
    ["3. Đường lây truyền bệnh SXH (C9)", "", ""],
    ["  Do muỗi chích (Đúng)", q9_c["Do muỗi chích"], f"{q9_c['Do muỗi chích']/N*100:.1f}%"],
    ["  Không biết", q9_c["Không biết"], f"{q9_c['Không biết']/N*100:.1f}%"],
    ["  Đường khác", q9_c["Khác (điền 9.1)"], f"{q9_c['Khác (điền 9.1)']/N*100:.1f}%"],
    ["  Khuyết", q9_c[None], f"{q9_c[None]/N*100:.1f}%"],
    ["4. Loại muỗi truyền bệnh SXH (C10)", "", ""],
    ["  Muỗi có vằn", q10_c["Muỗi có vằn"], f"{q10_c['Muỗi có vằn']/N*100:.1f}%"],
    ["  Muỗi sọc trắng đen", q10_c["Muỗi sọc trắng đen"], f"{q10_c['Muỗi sọc trắng đen']/N*100:.1f}%"],
    ["  Muỗi Aedes", q10_c["Muỗi Aedes"], f"{q10_c['Muỗi Aedes']/N*100:.1f}%"],
    ["  Muỗi khác", q10_c["Khác (điền 10.1)"], f"{q10_c['Khác (điền 10.1)']/N*100:.1f}%"],
    ["  Không nghe / Không biết", q10_c["Không nghe/Không biết"], f"{q10_c['Không nghe/Không biết']/N*100:.1f}%"],
    ["  Khuyết", q10_c[None], f"{q10_c[None]/N*100:.1f}%"],
    ["5. Thời điểm muỗi SXH thường chích (C11)", "", ""],
    ["  Ban ngày (Đúng)", q11_c["Ban ngày"], f"{q11_c['Ban ngày']/N*100:.1f}%"],
    ["  Ban đêm", q11_c["Ban đêm"], f"{q11_c['Ban đêm']/N*100:.1f}%"],
    ["  Ngày lẫn đêm", q11_c["Ngày lẫn đêm"], f"{q11_c['Ngày lẫn đêm']/N*100:.1f}%"],
    ["  Không biết", q11_c["Không biết"], f"{q11_c['Không biết']/N*100:.1f}%"],
    ["  Khuyết", q11_c[None], f"{q11_c[None]/N*100:.1f}%"],
    ["6. Nơi muỗi SXH đẻ trứng (C12 - Đa lựa chọn)", "", ""]])

for label, cnt, pct in get_multi_stat("12."):
    t2_rows.append([f"  {label}", cnt, f"{pct:.1f}%"])

t2_rows.extend([
    ["7. Cách phòng ngừa bệnh SXH (C13 - Đa lựa chọn)", "", ""]])

for label, cnt, pct in get_multi_stat("13."):
    if "chị hãy cho biết" not in label:
        t2_rows.append([f"  {label}", cnt, f"{pct:.1f}%"])

add_table_sheet("2. Kiến thức SXH", [
    ("BẢNG 2: KIẾN THỨC VỀ BỆNH SỐT XUẤT HUYẾT CỦA CỘNG ĐỒNG (N = 601)", ["Nội dung kiến thức", "Tần số (n)", "Tỷ lệ (%)"], t2_rows)
])

# ==================== TABLE 3: THỰC HÀNH PHÒNG CHỐNG SXH ====================
q14_c = Counter(col("14. Nhà anh chị có lu, khạp, phuy, hoặc hồ trữ nước để uống hay sinh hoạt hay không?  "))

# Q15
q15_raw = col("15. Tổng cộng, gia đình có bao nhiêu vật chứa nước?")
q15_cats = []
q15_nums = []
for v in q15_raw:
    if v is None: q15_cats.append("Không có / Khuyết")
    else:
        try:
            num = float(v)
            q15_nums.append(num)
            if num == 0: q15_cats.append("0 cái")
            elif 1 <= num <= 2: q15_cats.append("1 - 2 cái")
            elif 3 <= num <= 5: q15_cats.append("3 - 5 cái")
            else: q15_cats.append("≥ 6 cái")
        except: q15_cats.append("Khác")
q15_c = Counter(q15_cats)

q16_c = Counter(col("16. Anh chị có súc rửa vật chứa nuớc không?"))

# Q17
q17_raw = col("17. Nếu có, Bao lâu súc rửa 1 lần?")
q17_cats = []
for v in q17_raw:
    if v is None: q17_cats.append("Không áp dụng / Khuyết")
    else:
        try:
            num = float(v)
            if num < 7: q17_cats.append("< 7 ngày (< 1 tuần)")
            elif num == 7: q17_cats.append("7 ngày (1 tuần/lần)")
            elif 7 < num <= 14: q17_cats.append("8 - 14 ngày (1 - 2 tuần/lần)")
            elif 14 < num <= 30: q17_cats.append("15 - 30 ngày (2 tuần - 1 tháng/lần)")
            else: q17_cats.append("> 30 ngày (> 1 tháng/lần)")
        except: q17_cats.append("Khác")
q17_c = Counter(q17_cats)

q18_c = Counter(col("18. Các vật trữ nước có nắp đậy hay không?"))
q19_c = Counter(col("19. Có đậy nắp ngay sau khi sử dụng không?"))
q21_c = Counter(col("21. Anh chị có áp dụng biện pháp diệt lăng quăng hoặc biện pháp xua diệt muỗi nào không?"))
q23_c = Counter(col("23. Tủ đựng thức ăn của anh chị có chén chống kiến hay không? "))
q24_c = Counter(col("24. Nếu có, Anh chị đã làm gì để muỗi không đẻ trứng vào chén nước chống kiến?  "))
q25_c = Counter(col("25. Anh chị xử lý tô chén bể như thế nào?"))
q26_c = Counter(col("26. Ban ngày gia đình có ngủ mùng không? "))
q27_c = Counter(col("27. Nếu địa phương mở cuộc vận động phòng chống SXH, anh chị có tham gia không?"))

t3_rows = [
    ["1. Có vật trữ nước (lu, khạp, phuy, hồ) trong gia đình (C14)", "", ""],
    ["  Có", q14_c["Có"], f"{q14_c['Có']/N*100:.1f}%"],
    ["  Không", q14_c["Không"], f"{q14_c['Không']/N*100:.1f}%"],
    ["  Khuyết", q14_c[None], f"{q14_c[None]/N*100:.1f}%"],
    ["2. Số lượng vật chứa nước của gia đình (C15)", "", ""],
    ["  0 cái", q15_c["0 cái"], f"{q15_c['0 cái']/N*100:.1f}%"],
    ["  1 - 2 cái", q15_c["1 - 2 cái"], f"{q15_c['1 - 2 cái']/N*100:.1f}%"],
    ["  3 - 5 cái", q15_c["3 - 5 cái"], f"{q15_c['3 - 5 cái']/N*100:.1f}%"],
    ["  ≥ 6 cái", q15_c["≥ 6 cái"], f"{q15_c['≥ 6 cái']/N*100:.1f}%"],
    ["  Không có vật chứa / Khuyết", q15_c["Không có / Khuyết"], f"{q15_c['Không có / Khuyết']/N*100:.1f}%"],
    ["  Số vật chứa TB ± ĐLC (Min - Max)", f"{statistics.mean(q15_nums):.1f} ± {statistics.stdev(q15_nums):.1f}", f"Trung vị: {statistics.median(q15_nums):.1f} ({min(q15_nums):.0f} - {max(q15_nums):.0f})"],
    ["3. Có súc rửa vật chứa nước hay không (C16)", "", ""],
    ["  Có", q16_c["Có"], f"{q16_c['Có']/N*100:.1f}%"],
    ["  Không", q16_c["Không"], f"{q16_c['Không']/N*100:.1f}%"],
    ["  Khuyết", q16_c[None], f"{q16_c[None]/N*100:.1f}%"],
    ["4. Tần suất súc rửa vật chứa nước (C17)", "", ""],
    ["  < 7 ngày (< 1 tuần/lần)", q17_c["< 7 ngày (< 1 tuần)"], f"{q17_c['< 7 ngày (< 1 tuần)']/N*100:.1f}%"],
    ["  7 ngày (1 tuần/lần)", q17_c["7 ngày (1 tuần/lần)"], f"{q17_c['7 ngày (1 tuần/lần)']/N*100:.1f}%"],
    ["  8 - 14 ngày (1 - 2 tuần/lần)", q17_c["8 - 14 ngày (1 - 2 tuần/lần)"], f"{q17_c['8 - 14 ngày (1 - 2 tuần/lần)']/N*100:.1f}%"],
    ["  15 - 30 ngày (2 tuần - 1 tháng/lần)", q17_c["15 - 30 ngày (2 tuần - 1 tháng/lần)"], f"{q17_c['15 - 30 ngày (2 tuần - 1 tháng/lần)']/N*100:.1f}%"],
    ["  > 30 ngày (> 1 tháng/lần)", q17_c["> 30 ngày (> 1 tháng/lần)"], f"{q17_c['> 30 ngày (> 1 tháng/lần)']/N*100:.1f}%"],
    ["  Không áp dụng / Khuyết", q17_c["Không áp dụng / Khuyết"], f"{q17_c['Không áp dụng / Khuyết']/N*100:.1f}%"],
    ["5. Nắp đậy các vật chứa nước (C18)", "", ""],
    ["  Tất cả có nắp đậy", q18_c["Tất cả có nắp đậy"], f"{q18_c['Tất cả có nắp đậy']/N*100:.1f}%"],
    ["  Chỉ một số cái có nắp đậy", q18_c["Chỉ một số cái có nắp đậy"], f"{q18_c['Chỉ một số cái có nắp đậy']/N*100:.1f}%"],
    ["  Không có nắp đậy", q18_c["Không"], f"{q18_c['Không']/N*100:.1f}%"],
    ["  Khuyết", q18_c[None], f"{q18_c[None]/N*100:.1f}%"],
    ["6. Đậy nắp ngay sau khi sử dụng (C19)", "", ""],
    ["  Có", q19_c["Có"], f"{q19_c['Có']/N*100:.1f}%"],
    ["  Không", q19_c["Không"], f"{q19_c['Không']/N*100:.1f}%"],
    ["  Khuyết", q19_c[None], f"{q19_c[None]/N*100:.1f}%"],
    ["7. Lý do đậy nắp vật trữ nước (C20 - Đa lựa chọn)", "", ""]]

for label, cnt, pct in get_multi_stat("20."):
    t3_rows.append([f"  {label}", cnt, f"{pct:.1f}%"])

t3_rows.extend([
    ["8. Áp dụng biện pháp diệt lăng quăng / xua muỗi (C21)", "", ""],
    ["  Có", q21_c["Có"], f"{q21_c['Có']/N*100:.1f}%"],
    ["  Không", q21_c["Không"], f"{q21_c['Không']/N*100:.1f}%"],
    ["  Khuyết", q21_c[None], f"{q21_c[None]/N*100:.1f}%"],
    ["9. Biện pháp loại trừ nơi muỗi đẻ trứng (C22.1 - Đa lựa chọn)", "", ""]])

for label, cnt, pct in get_multi_stat("22.1."):
    t3_rows.append([f"  {label}", cnt, f"{pct:.1f}%"])

t3_rows.extend([
    ["10. Biện pháp phòng muỗi cắn và diệt muỗi (C22.2 - Đa lựa chọn)", "", ""]])

for label, cnt, pct in get_multi_stat("22.2."):
    t3_rows.append([f"  {label}", cnt, f"{pct:.1f}%"])

t3_rows.extend([
    ["11. Biện pháp loại bỏ nơi muỗi trú ẩn (C22.3 - Đa lựa chọn)", "", ""]])

for label, cnt, pct in get_multi_stat("22.3."):
    t3_rows.append([f"  {label}", cnt, f"{pct:.1f}%"])

t3_rows.extend([
    ["12. Chén chống kiến tủ thức ăn (C23)", "", ""],
    ["  Không có", q23_c["Không"], f"{q23_c['Không']/N*100:.1f}%"],
    ["  Có chén chống kiến", q23_c["Có"], f"{q23_c['Có']/N*100:.1f}%"],
    ["  Khuyết", q23_c[None], f"{q23_c[None]/N*100:.1f}%"],
    ["13. Biện pháp xử lý chén chống kiến (C24)", "", ""],
    ["  Không làm gì", q24_c["Không làm gì"], f"{q24_c['Không làm gì']/N*100:.1f}%"],
    ["  Bỏ muối vào chén", q24_c["Bỏ muối vào chén"], f"{q24_c['Bỏ muối vào chén']/N*100:.1f}%"],
    ["  Đổ dầu (nhớt) vào chén", q24_c["Đổ dầu (nhớt) vào chén"], f"{q24_c['Đổ dầu (nhớt) vào chén']/N*100:.1f}%"],
    ["  Cách làm khác", q24_c["Khác (điền 24.1)"], f"{q24_c['Khác (điền 24.1)']/N*100:.1f}%"],
    ["  Không áp dụng / Khuyết", q24_c[None], f"{q24_c[None]/N*100:.1f}%"],
    ["14. Xử lý tô chén bể quanh nhà (C25)", "", ""],
    ["  Bỏ thùng rác (Đúng)", q25_c["Bỏ thùng rác"], f"{q25_c['Bỏ thùng rác']/N*100:.1f}%"],
    ["  Chôn hoặc úp xuống (Đúng)", q25_c["Chôn hoặc úp xuống"], f"{q25_c['Chôn hoặc úp xuống']/N*100:.1f}%"],
    ["  Vứt bừa ra vườn", q25_c["Vứt bừa ra vườn"], f"{q25_c['Vứt bừa ra vườn']/N*100:.1f}%"],
    ["  Cách làm khác", q25_c["Khác (điền 25.1)"], f"{q25_c['Khác (điền 25.1)']/N*100:.1f}%"],
    ["  Khuyết", q25_c[None], f"{q25_c[None]/N*100:.1f}%"],
    ["15. Thói quen ngủ mùng ban ngày (C26)", "", ""],
    ["  Không", q26_c["Không"], f"{q26_c['Không']/N*100:.1f}%"],
    ["  Có (Đúng)", q26_c["Có"], f"{q26_c['Có']/N*100:.1f}%"],
    ["  Khuyết", q26_c[None], f"{q26_c[None]/N*100:.1f}%"],
    ["16. Sẵn sàng tham gia cuộc vận động phòng chống SXH (C27)", "", ""],
    ["  Có", q27_c["Có"], f"{q27_c['Có']/N*100:.1f}%"],
    ["  Không", q27_c["Không"], f"{q27_c['Không']/N*100:.1f}%"],
    ["  Không ý kiến", q27_c["Không ý kiến"], f"{q27_c['Không ý kiến']/N*100:.1f}%"],
    ["  Khác / Khuyết", q27_c[None] + q27_c.get("Có Không", 0) + q27_c.get("Có Không ý kiến", 0) + q27_c.get("Không Không ý kiến", 0), f"{(q27_c[None] + q27_c.get('Có Không', 0) + q27_c.get('Có Không ý kiến', 0) + q27_c.get('Không Không ý kiến', 0))/N*100:.1f}%"]
])

add_table_sheet("3. Thực hành SXH", [
    ("BẢNG 3: THỰC HÀNH PHÒNG CHỐNG SỐT XUẤT HUYẾT CỦA CỘNG ĐỒNG (N = 601)", ["Nội dung thực hành", "Tần số (n)", "Tỷ lệ (%)"], t3_rows)
])

# Save
out_wb.remove(default_sheet)
out_file = r"C:\Users\ThinkPad\Downloads\chapteros11\Bao_cao_Thong_ke_SXH_Cong_dong.xlsx"
out_wb.save(out_file)
print("SUCCESS: File generated at", out_file)
